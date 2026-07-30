"""
FastAPI backend for the Multi-Agent Data Analyst.

Exposes REST endpoints that the React frontend uses:
  - POST /api/upload    → Upload a CSV, parse schema, store in session
  - POST /api/query     → Run the LangGraph pipeline on a user query
  - GET  /api/health    → Health check
"""

import io
import json
import uuid
import logging
import traceback
import re

import pandas as pd
import plotly
import plotly.io as pio
import plotly.graph_objects as go
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse

from orchestrator.nodes import data_understanding
from orchestrator.graph import build_graph
from db.db import log_query, log_retry

# -- Logging ------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(name)-30s  %(levelname)-7s  %(message)s",
)
logger = logging.getLogger(__name__)

# -- Constants ----------------------------------------------------------------
MAX_UPLOAD_MB = 200
MAX_RETRIES = 3

# -- App ----------------------------------------------------------------------
app = FastAPI(
    title="Multi-Agent Data Analyst API",
    version="1.0.0",
)

# CORS for React dev server
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
async def startup_event():
    logger.info("Warming up Kaleido engine for fast Word exports...")
    try:
        # Create a tiny dummy figure
        fig = go.Figure(go.Scatter(x=[1], y=[1]))
        # Run to_image to initialize the Kaleido process
        pio.to_image(fig, format="png")
        logger.info("Kaleido engine warmed up successfully.")
    except Exception as e:
        logger.warning(f"Could not warm up Kaleido engine: {e}")

# -- In-memory session store (keyed by session_id) ----------------------------
# In production this would be Redis / DB backed.
_sessions: dict[str, dict] = {}


def _clean_json_val(val):
    if pd.isna(val):
        return None
    if isinstance(val, (pd.Timestamp, pd.Timedelta)):
        return str(val)
    if hasattr(val, "item"):
        return val.item()
    return val


def _serialize_result(execution_result, output_type: str) -> dict:
    """Convert execution_result to JSON-safe structures."""
    if execution_result is None:
        return {"type": "none", "data": None}

    # Plotly figure
    if hasattr(execution_result, "to_json"):
        try:
            return {
                "type": "plotly",
                "data": json.loads(plotly.io.to_json(execution_result)),
            }
        except Exception:
            pass

    # pandas DataFrame
    if isinstance(execution_result, pd.DataFrame):
        df_sub = execution_result.head(200)
        cols = [str(c) for c in df_sub.columns]
        raw_records = df_sub.to_dict(orient="records")
        clean_rows = [
            {str(k): _clean_json_val(v) for k, v in row.items()}
            for row in raw_records
        ]
        return {
            "type": "table",
            "data": {
                "columns": cols,
                "rows": clean_rows,
            },
        }

    # pandas Series
    if isinstance(execution_result, pd.Series):
        df_sub = execution_result.head(200).to_frame()
        cols = [str(c) for c in df_sub.columns]
        raw_records = df_sub.to_dict(orient="records")
        clean_rows = [
            {str(k): _clean_json_val(v) for k, v in row.items()}
            for row in raw_records
        ]
        return {
            "type": "table",
            "data": {
                "columns": cols,
                "rows": clean_rows,
            },
        }

    # Scalar / primitive
    return {"type": "scalar", "data": str(execution_result)}


# -- Routes -------------------------------------------------------------------

@app.get("/api/health")
async def health():
    return {"status": "ok", "sessions": len(_sessions)}


@app.post("/api/upload")
async def upload_csv(file: UploadFile = File(...)):
    """Upload a CSV, parse it, and return a session_id + schema summary."""
    if not file.filename.endswith(".csv"):
        raise HTTPException(status_code=400, detail="Only CSV files are accepted.")

    raw_bytes = await file.read()
    size_mb = len(raw_bytes) / (1024 * 1024)
    if size_mb > MAX_UPLOAD_MB:
        raise HTTPException(
            status_code=413,
            detail=f"File too large ({size_mb:.1f} MB). Maximum is {MAX_UPLOAD_MB} MB.",
        )

    # Create a file-like object for data_understanding
    class _FakeUpload:
        def __init__(self, data: bytes, name: str):
            self._buf = io.BytesIO(data)
            self.name = name
            self.size = len(data)
        def read(self):
            return self._buf.read()
        def seek(self, pos):
            self._buf.seek(pos)

    fake = _FakeUpload(raw_bytes, file.filename)

    try:
        result = data_understanding.run({}, fake)
    except Exception as exc:
        logger.error("Failed to parse CSV: %s", exc)
        raise HTTPException(status_code=422, detail=f"Failed to parse CSV: {exc}")

    df: pd.DataFrame = result["df"]
    schema_context: str = result["schema_context"]

    session_id = str(uuid.uuid4())
    _sessions[session_id] = {
        "df": df,
        "schema_context": schema_context,
        "filename": file.filename,
    }

    return {
        "session_id": session_id,
        "filename": file.filename,
        "rows": int(df.shape[0]),
        "columns": int(df.shape[1]),
        "column_names": list(df.columns),
        "column_types": {col: str(dtype) for col, dtype in zip(df.columns, df.dtypes)},
        "schema_context": schema_context,
    }


@app.post("/api/query")
async def run_query(session_id: str = Form(...), query: str = Form(...)):
    """Run a user query through the LangGraph multi-agent pipeline."""
    session = _sessions.get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Session not found. Please upload a CSV first.")

    initial_state = {
        "df": session["df"],
        "schema_context": session["schema_context"],
        "user_query": query,
        "output_type": "retrieval",
        "analysis_plan": "",
        "code_type": "pandas",
        "generated_code": "",
        "execution_result": None,
        "error_message": None,
        "failed_attempts": [],
        "retry_count": 0,
        "max_retries": MAX_RETRIES,
        "insight_text": None,
        "fallback_used": False,
    }

    try:
        graph = build_graph()
        result_state = graph.invoke(initial_state)
    except Exception as exc:
        logger.error("Pipeline error: %s\n%s", exc, traceback.format_exc())
        raise HTTPException(status_code=500, detail=f"Pipeline error: {exc}")

    # Log to DB
    try:
        query_id = log_query(
            user_query=query,
            output_type=result_state.get("output_type"),
            analysis_plan=result_state.get("analysis_plan"),
            code_type=result_state.get("code_type"),
            generated_code=result_state.get("generated_code"),
            execution_success=result_state.get("error_message") is None,
            retry_count=result_state.get("retry_count", 0),
            fallback_used=result_state.get("fallback_used", False),
            insight_text=result_state.get("insight_text"),
            error_message=result_state.get("error_message"),
        )
        for i, attempt in enumerate(result_state.get("failed_attempts", []), 1):
            log_retry(query_id, i, attempt["code"], attempt["error"])
    except Exception as exc:
        logger.error("Failed to log query: %s", exc)

    # Serialize result
    serialized = _serialize_result(
        result_state.get("execution_result"),
        result_state.get("output_type", "retrieval"),
    )

    return {
        "output_type": result_state.get("output_type"),
        "code_type": result_state.get("code_type"),
        "generated_code": result_state.get("generated_code"),
        "analysis_plan": result_state.get("analysis_plan"),
        "execution_result": serialized,
        "insight_text": result_state.get("insight_text"),
        "retry_count": result_state.get("retry_count", 0),
        "fallback_used": result_state.get("fallback_used", False),
        "failed_attempts": [
            {"code": a.get("code", ""), "error": a.get("error", "")}
            for a in result_state.get("failed_attempts", [])
        ],
    }


@app.post("/api/export_report")
async def export_report(
    filename: str = Form("session"),
    history_json: str = Form(...)
):
    from docx.oxml.ns import qn

    history = json.loads(history_json)

    safe_filename = filename.replace(".csv", "")
    safe_filename = re.sub(r'[^a-zA-Z0-9_-]', '_', safe_filename)

    doc = Document()

    # -- Styles ---------------------------------------------------------------
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(10)
    font.color.rgb = RGBColor(0x37, 0x41, 0x51)

    # -- Title ----------------------------------------------------------------
    title = doc.add_heading('Data Analysis Report', level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    doc.add_paragraph(f'Dataset: {filename}')
    doc.add_paragraph(f'Generated: {pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")}')
    doc.add_paragraph('').paragraph_format.space_after = Pt(6)

    # -- Each query entry -----------------------------------------------------
    for idx, entry in enumerate(history):
        query = entry.get("query", "")
        heading = doc.add_heading(f'Query {idx + 1}: {query}', level=2)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

        # Analysis / Insight
        insight = entry.get("insight_text")
        if insight:
            doc.add_heading('Analysis', level=3)
            for line in insight.split('\n'):
                if line.strip():
                    doc.add_paragraph(line.strip())

        # Execution result
        result = entry.get("execution_result")
        if result and result.get("type") == "plotly" and result.get("data"):
            # Render Plotly figure to PNG and embed in the document
            doc.add_heading('Chart', level=3)
            try:
                fig_dict = result["data"]
                # Set white background for better print appearance
                if "layout" not in fig_dict:
                    fig_dict["layout"] = {}
                fig_dict["layout"]["paper_bgcolor"] = "white"
                fig_dict["layout"]["plot_bgcolor"] = "white"

                # Automatically show legend for pie charts and multi-trace figures
                showlegend = fig_dict["layout"].get("showlegend")
                if showlegend is None:
                    traces = fig_dict.get("data", [])
                    is_pie_or_donut = any(t.get("type") == "pie" for t in traces)
                    has_multiple_traces = len(traces) > 1
                    if is_pie_or_donut or has_multiple_traces:
                        showlegend = True

                if showlegend:
                    fig_dict["layout"]["showlegend"] = True
                    if "margin" not in fig_dict["layout"]:
                        fig_dict["layout"]["margin"] = {}
                    fig_dict["layout"]["margin"]["r"] = 100

                img_bytes = pio.to_image(fig_dict, format="png", width=900, height=500, scale=1.5)
                img_stream = io.BytesIO(img_bytes)
                doc.add_picture(img_stream, width=Inches(6))
                # Center the image
                last_paragraph = doc.paragraphs[-1]
                last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as chart_err:
                logger.warning(f"Failed to render chart for export: {chart_err}")
                doc.add_paragraph(f"[Chart could not be rendered: {chart_err}]")

        elif result and result.get("type") == "table" and result.get("data"):
            raw = result["data"]
            rows = raw.get("rows", []) if isinstance(raw, dict) else raw
            if rows and len(rows) > 0:
                doc.add_heading('Result Data', level=3)
                cols = list(rows[0].keys())
                display_rows = rows[:15]
                table = doc.add_table(rows=1 + len(display_rows), cols=len(cols))
                table.style = 'Table Grid'

                # Header row
                for ci, col_name in enumerate(cols):
                    cell = table.rows[0].cells[ci]
                    cell.text = str(col_name)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                            run.font.size = Pt(9)

                # Data rows
                for ri, row in enumerate(display_rows):
                    for ci, c in enumerate(cols):
                        val = row.get(c)
                        if val is None:
                            val = "—"
                        cell = table.rows[ri + 1].cells[ci]
                        cell.text = str(val)
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(9)

                if len(rows) > 15:
                    doc.add_paragraph(f'Showing top 15 of {len(rows)} rows.').italic = True

        elif result and result.get("type") == "scalar":
            doc.add_heading('Result', level=3)
            doc.add_paragraph(str(result.get('data', '')))

        # Generated code
        code = entry.get("generated_code")
        if code:
            doc.add_heading('Python Code', level=3)
            code_para = doc.add_paragraph()
            code_run = code_para.add_run(code)
            code_run.font.name = 'Courier New'
            code_run.font.size = Pt(8.5)
            code_run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            # Set monospace font for the run element
            rPr = code_run._element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn('w:ascii'), 'Courier New')
            rFonts.set(qn('w:hAnsi'), 'Courier New')

        # Separator between queries
        doc.add_paragraph('').paragraph_format.space_after = Pt(4)

    # -- Serialize to bytes ---------------------------------------------------
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    docx_bytes = buffer.getvalue()

    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="data_analysis_report_{safe_filename}.docx"'
        }
    )
